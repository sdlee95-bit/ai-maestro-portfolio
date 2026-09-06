# -*- coding: utf-8 -*-
"""
make_assets.py — assets/ 재생성

원본 자료가 있을 때만 실행하면 된다 (평상시엔 불필요).
  · 2026 업무보고 서식 1부.pptx   → base.pptx, campus*.png
  · 부산대로고 BASIC-01-1.ai      → pnu_emblem*.png

실행 : python make_assets.py <원본폴더>
필요 : python-pptx, pillow, pymupdf
"""
import os
import re
import shutil
import sys
import zipfile

import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation

SRC = sys.argv[1] if len(sys.argv) > 1 else "."
HERE = os.path.dirname(os.path.abspath(__file__))
A = os.path.normpath(os.path.join(HERE, "..", "assets"))
os.makedirs(A, exist_ok=True)

PPTX = os.path.join(SRC, "2026 업무보고 서식 1부.pptx")
AI = os.path.join(SRC, "부산대로고 BASIC-01-1.ai")

# ── 1. 부산대 엠블럼 : .ai(PDF 1.4) → 고해상도 투명 PNG ──────────────
doc = fitz.open(AI)
pix = doc[0].get_pixmap(matrix=fitz.Matrix(8, 8), alpha=True)
pix.save(os.path.join(A, "_logo_raw.png"))
img = Image.open(os.path.join(A, "_logo_raw.png")).convert("RGBA")
img = img.crop(img.split()[-1].getbbox())          # 알파 기준 여백 제거
img.save(os.path.join(A, "pnu_emblem.png"))

alpha = img.split()[-1]                             # 화이트 모노 버전
white = Image.merge("RGBA", (Image.new("L", img.size, 255),
                             Image.new("L", img.size, 255),
                             Image.new("L", img.size, 255), alpha))
white.save(os.path.join(A, "pnu_emblem_white.png"))
os.remove(os.path.join(A, "_logo_raw.png"))

# ── 2. 캠퍼스 사진 : 원본 pptx의 media 에서 추출 ────────────────────
with zipfile.ZipFile(PPTX) as z:
    with open(os.path.join(A, "campus.png"), "wb") as f:      # 컬러 (표지)
        f.write(z.read("ppt/media/image1.png"))
    with open(os.path.join(A, "campus_gray.png"), "wb") as f:  # 흑백 (목차 띠)
        f.write(z.read("ppt/media/image3.png"))

# ── 3. base.pptx : 슬라이드만 제거, 테마·임베드 서체는 보존 ─────────
prs = Presentation(PPTX)
lst = prs.slides._sldIdLst
for sld in list(lst):
    prs.part.drop_rel(sld.rId)
    lst.remove(sld)
tmp = os.path.join(A, "_base_tmp.pptx")
prs.save(tmp)

# ── 4. 맑은 고딕 임베드 서체 제거 ───────────────────────────────────
# 원본 서식에 딸려온 것으로 이 라이브러리는 사용하지 않는다.
# Microsoft 독점 서체라 재배포가 불가하고, 14.1MB 중 12.4MB를 차지한다.
# 제거해도 페이퍼로지·공체는 임베드 상태로 남아 렌더링에 영향이 없다.
DROP_FACE = "맑은 고딕"
zin = zipfile.ZipFile(tmp)
pres_xml = zin.read("ppt/presentation.xml").decode("utf-8")
rels_xml = zin.read("ppt/_rels/presentation.xml.rels").decode("utf-8")

drop_rids = []


def _kill(m):
    if f'typeface="{DROP_FACE}"' in m.group(0):
        drop_rids.extend(re.findall(r'r:id="(rId\d+)"', m.group(0)))
        return ""
    return m.group(0)


pres_xml = re.sub(r"<p:embeddedFont>.*?</p:embeddedFont>", _kill, pres_xml, flags=re.S)

drop_parts = []
for rid in drop_rids:
    m = re.search(r'<Relationship Id="%s"[^>]*Target="([^"]+)"[^>]*/>' % rid, rels_xml)
    drop_parts.append("ppt/" + m.group(1).lstrip("/"))
    rels_xml = rels_xml.replace(m.group(0), "")

with zipfile.ZipFile(os.path.join(A, "base.pptx"), "w", zipfile.ZIP_DEFLATED) as zo:
    for item in zin.infolist():
        if item.filename in drop_parts:
            continue
        data = zin.read(item.filename)
        if item.filename == "ppt/presentation.xml":
            data = pres_xml.encode("utf-8")
        elif item.filename == "ppt/_rels/presentation.xml.rels":
            data = rels_xml.encode("utf-8")
        zo.writestr(item, data)
zin.close()
os.remove(tmp)
print(f"제거한 임베드 서체: {DROP_FACE} ({len(drop_parts)}개 파트)")

for f in sorted(os.listdir(A)):
    print(f"{f:26s} {os.path.getsize(os.path.join(A, f)):>10,d} bytes")
print("남은 슬라이드:", len(Presentation(os.path.join(A, "base.pptx")).slides))
